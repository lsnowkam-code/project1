#!/usr/bin/env python3
"""
Создание тестового Word шаблона для проверки динамического маппинга
"""

from docx import Document
from docx.shared import Inches

def create_test_word_template():
    """Создает тестовый Word шаблон с таблицей внеоборотных активов"""

    doc = Document()

    # Заголовок документа
    doc.add_heading('Бюллетень 17.2.8 - Тестовый шаблон', 0)

    # Таблица 3: Внеоборотные активы
    doc.add_paragraph('Таблица №3 - Внеоборотные активы организаций по видам экономической деятельности')
    doc.add_paragraph('')  # Пустая строка

    # Создаем таблицу
    table = doc.add_table(rows=5, cols=4)  # 4 колонки: OKVED, Наименование, 2022, 2023
    table.style = 'Table Grid'

    # Заголовки
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Код'
    hdr_cells[1].text = 'Наименование'
    hdr_cells[2].text = '2022'
    hdr_cells[3].text = '2023'

    # Данные (тестовые OKVED)
    test_data = [
        ('101.АГ', 'Растениеводство и животноводство'),
        ('102.АГ', 'Лесоводство и лесозаготовки'),
        ('103.АГ', 'Рыболовство и рыбоводство')
    ]

    for i, (okved, name) in enumerate(test_data, 1):
        row_cells = table.rows[i].cells
        row_cells[0].text = okved
        row_cells[1].text = name
        # Плейсхолдеры для данных - оставляем пустыми, генератор заполнит
        row_cells[2].text = ''  # DolgFinVl 2022
        row_cells[3].text = ''  # DolgFinVl 2023

    # Сохраняем шаблон
    template_path = 'Бюллетень_17.2.8_ШАБЛОН.docx'
    doc.save(template_path)
    print(f"✅ Тестовый шаблон создан: {template_path}")

    return template_path

if __name__ == "__main__":
    create_test_word_template()