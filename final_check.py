from docx import Document
from pathlib import Path

# Два файла для проверки
generated_correct = Path('/workspaces/project/output/Бюллетень_ПРОВЕРКА_ПРАВИЛЬНО.docx')
reference = Path('/workspaces/project/output/как должно быть!.docx')

if generated_correct.exists():
    print("✅ Файл сгенерирован!")
    
    gen = Document(generated_correct)
    ref = Document(reference)
    
    # Таблица 17
    t17_gen = gen.tables[16]
    t17_ref = ref.tables[16]
    
    print()
    print("=" * 80)
    print("ТАБЛИЦА 17 - ПРОВЕРКА СТРУКТУРЫ")
    print("=" * 80)
    print(f"Сгенерирована: {len(t17_gen.rows)} строк x {len(t17_gen.rows[0].cells)} столбцов")
    print(f"Эталон:       {len(t17_ref.rows)} строк x {len(t17_ref.rows[0].cells)} столбцов")
    
    if len(t17_gen.rows[0].cells) == len(t17_ref.rows[0].cells):
        print("✅ СТРУКТУРА СОВПАДАЕТ!")
    else:
        print("❌ СТРУКТУРА НЕ СОВПАДАЕТ!")
    
    # Подсчет заполненных ячеек
    def count_filled(table):
        filled = 0
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text and text != '—':
                    filled += 1
        return filled
    
    gen_filled = count_filled(t17_gen)
    ref_filled = count_filled(t17_ref)
    
    print()
    print(f"Сгенерирована: {gen_filled} ячеек заполнено")
    print(f"Эталон:       {ref_filled} ячеек заполнено")
    
    if gen_filled == ref_filled:
        print("✅ ЗАПОЛНЕНИЕ СОВПАДАЕТ!")
    else:
        print(f"❌ ЗАПОЛНЕНИЕ НЕ СОВПАДАЕТ ({abs(gen_filled - ref_filled)} ячеек разницы)")
    
else:
    print("⏳ Файл еще генерируется...")
