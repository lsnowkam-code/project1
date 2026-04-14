#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Скрипт для сравнения двух Word документов
Сравнивает сгенерированный документ с эталонным
"""

from docx import Document
from pathlib import Path
import pandas as pd

def extract_table_data(doc_path):
    """Извлекает данные из всех таблиц документа"""
    doc = Document(doc_path)
    all_tables_data = []

    for table_idx, table in enumerate(doc.tables):
        table_data = []
        for row_idx, row in enumerate(table.rows):
            row_data = []
            for cell_idx, cell in enumerate(row.cells):
                # Получаем текст из всех параграфов в ячейке
                cell_text = ' '.join([p.text.strip() for p in cell.paragraphs if p.text.strip()])
                row_data.append(cell_text)
            table_data.append(row_data)
        all_tables_data.append(table_data)

    return all_tables_data

def compare_documents(generated_path, reference_path):
    """Сравнивает два документа и возвращает различия"""

    print(f"📄 Сравниваем документы:")
    print(f"   Сгенерированный: {generated_path}")
    print(f"   Эталонный: {reference_path}")
    print()

    try:
        generated_data = extract_table_data(generated_path)
        reference_data = extract_table_data(reference_path)
    except Exception as e:
        print(f"❌ Ошибка при чтении документов: {e}")
        return None

    differences = []
    max_tables = max(len(generated_data), len(reference_data))

    for table_idx in range(max_tables):
        gen_table = generated_data[table_idx] if table_idx < len(generated_data) else []
        ref_table = reference_data[table_idx] if table_idx < len(reference_data) else []

        max_rows = max(len(gen_table), len(ref_table))

        for row_idx in range(max_rows):
            gen_row = gen_table[row_idx] if row_idx < len(gen_table) else []
            ref_row = ref_table[row_idx] if row_idx < len(ref_table) else []

            max_cols = max(len(gen_row), len(ref_row))

            for col_idx in range(max_cols):
                gen_cell = gen_row[col_idx] if col_idx < len(gen_row) else ""
                ref_cell = ref_row[col_idx] if col_idx < len(ref_row) else ""

                # Очищаем от лишних пробелов и сравниваем
                gen_clean = gen_cell.strip()
                ref_clean = ref_cell.strip()

                if gen_clean != ref_clean:
                    differences.append({
                        'Таблица': table_idx + 1,
                        'Строка': row_idx + 1,
                        'Столбец': col_idx + 1,
                        'Сгенерировано': gen_clean,
                        'Эталон': ref_clean,
                        'Совпадение': gen_clean == ref_clean
                    })

    return differences

def main():
    base_dir = Path(__file__).parent

    # Пути к файлам
    generated_file = base_dir / "output" / "Бюллетень_17.2.8_ГОТОВЫЙ.docx"
    reference_file = base_dir / "output" / "как должно быть!.docx"

    # Альтернативный путь к эталонному файлу
    if not reference_file.exists():
        reference_file = base_dir / "input" / "Как должно быть.docx"

    if not generated_file.exists():
        print(f"❌ Сгенерированный файл не найден: {generated_file}")
        return

    if not reference_file.exists():
        print(f"❌ Эталонный файл не найден: {reference_file}")
        return

    # Сравниваем документы
    differences = compare_documents(generated_file, reference_file)

    if differences is None:
        return

    print(f"🔍 Найдено различий: {len(differences)}")
    print()

    if differences:
        # Сохраняем различия в Excel
        df = pd.DataFrame(differences)
        output_file = base_dir / "output" / "сравнение_документов.xlsx"
        df.to_excel(output_file, index=False)
        print(f"📊 Подробные различия сохранены в: {output_file}")
        print()

        # Показываем первые 10 различий
        print("📋 Первые 10 различий:")
        print("-" * 80)
        for i, diff in enumerate(differences[:10]):
            print(f"{i+1}. Таблица {diff['Таблица']}, строка {diff['Строка']}, столбец {diff['Столбец']}")
            print(f"   Сгенерировано: '{diff['Сгенерировано']}'")
            print(f"   Эталон:        '{diff['Эталон']}'")
            print()

        if len(differences) > 10:
            print(f"... и ещё {len(differences) - 10} различий")
    else:
        print("✅ ДОКУМЕНТЫ ПОЛНОСТЬЮ СОВПАДАЮТ! 🎉")

if __name__ == "__main__":
    main()