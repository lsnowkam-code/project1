# compare_results.py - Сравнение результатов с эталоном
"""
Сравнивает сгенерированный документ с эталоном
"""

from docx import Document
import difflib
from pathlib import Path


def extract_text_from_docx(docx_path):
    """Извлекает текст из Word документа"""
    doc = Document(docx_path)
    text = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())

    # Также извлекаем текст из таблиц
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    text.append(cell_text)

    return '\n'.join(text)


def compare_documents(generated_path, reference_path):
    """Сравнивает два документа и показывает различия"""

    print("📄 Извлечение текста из документов...")

    try:
        generated_text = extract_text_from_docx(generated_path)
        reference_text = extract_text_from_docx(reference_path)

        print(f"📏 Длина сгенерированного текста: {len(generated_text)} символов")
        print(f"📏 Длина эталонного текста: {len(reference_text)} символов")

        if generated_text == reference_text:
            print("\n🎉 ДОКУМЕНТЫ ИДЕНТИЧНЫ! 100% совпадение!")
            return True

        # Показываем различия
        print("\n🔍 НАЙДЕНЫ РАЗЛИЧИЯ:")
        print("=" * 50)

        # Разбиваем на строки для построчного сравнения
        generated_lines = generated_text.split('\n')
        reference_lines = reference_text.split('\n')

        # Используем difflib для красивого сравнения
        diff = list(difflib.unified_diff(
            reference_lines,
            generated_lines,
            fromfile='ЭТАЛОН',
            tofile='СГЕНЕРИРОВАННЫЙ',
            lineterm='',
            n=3  # контекст
        ))

        if diff:
            print("Различия (первые 50 строк):")
            for line in diff[:50]:
                if line.startswith('+'):
                    print(f"🟢 {line}")
                elif line.startswith('-'):
                    print(f"🔴 {line}")
                elif line.startswith('@'):
                    print(f"📍 {line}")
                else:
                    print(f"   {line}")

            if len(diff) > 50:
                print(f"\n... и еще {len(diff) - 50} различий")
        else:
            print("Странно, тексты разные, но unified_diff не показал различий...")

        # Статистика
        similarity = difflib.SequenceMatcher(None, generated_text, reference_text).ratio()
        print(f"📈 Процент сходства: {similarity:.1%}")

        return False

    except Exception as e:
        print(f"❌ Ошибка при сравнении: {e}")
        return False


def main():
    output_dir = Path("/workspaces/project/output")

    generated_file = output_dir / "Бюллетень_17.2.8_ГОТОВЫЙ.docx"
    reference_file = output_dir / "как должно быть!.docx"

    print("🔍 СРАВНЕНИЕ РЕЗУЛЬТАТОВ С ЭТАЛОНОМ")
    print("=" * 50)
    print(f"📄 Сгенерированный: {generated_file.name}")
    print(f"🎯 Эталон: {reference_file.name}")
    print()

    if not generated_file.exists():
        print(f"❌ Сгенерированный файл не найден: {generated_file}")
        return

    if not reference_file.exists():
        print(f"❌ Эталонный файл не найден: {reference_file}")
        return

    # Сравниваем
    is_identical = compare_documents(generated_file, reference_file)

    print("\n" + "=" * 50)
    if is_identical:
        print("✅ РЕЗУЛЬТАТ: ДОСТИГНУТА ЦЕЛЬ!")
        print("   Сгенерированный документ идентичен эталону!")
    else:
        print("⚠️ РЕЗУЛЬТАТ: ЕСТЬ РАЗЛИЧИЯ")
        print("   Нужно проанализировать и исправить")
    print("=" * 50)


if __name__ == "__main__":
    main()
