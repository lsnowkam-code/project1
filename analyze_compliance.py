from docx import Document
from pathlib import Path

def analyze_documents():
    """Детальный анализ соответствия двух документов"""
    
    generated = Document('/workspaces/project/output/Бюллетень_17.2.8_ГОТОВЫЙ.docx')
    reference = Document('/workspaces/project/output/как должно быть!.docx')
    
    print("=" * 100)
    print("АНАЛИЗ СООТВЕТСТВИЯ: Бюллетень_17.2.8_ГОТОВЫЙ.docx vs как должно быть!.docx")
    print("=" * 100)
    print()
    
    # Базовая информация
    print(f"Таблиц в сгенерированном: {len(generated.tables)}")
    print(f"Таблиц в эталоне: {len(reference.tables)}")
    print()
    
    def count_filled(table):
        filled = 0
        total = 0
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                total += 1
                if text and text != '—':
                    filled += 1
        return filled, total
    
    # Анализ каждой таблицы
    differences = []
    
    for idx in range(min(len(generated.tables), len(reference.tables))):
        gen_table = generated.tables[idx]
        ref_table = reference.tables[idx]
        
        gen_filled, gen_total = count_filled(gen_table)
        ref_filled, ref_total = count_filled(ref_table)
        
        gen_pct = (gen_filled / gen_total * 100) if gen_total > 0 else 0
        ref_pct = (ref_filled / ref_total * 100) if ref_total > 0 else 0
        
        # Определяем разницу
        struct_match = (len(gen_table.rows) == len(ref_table.rows)) and \
                       (len(gen_table.rows[0].cells) == len(ref_table.rows[0].cells))
        
        fill_match = abs(gen_pct - ref_pct) < 0.1  # Допуск 0.1%
        
        if not struct_match or not fill_match:
            differences.append({
                'table': idx + 1,
                'gen_rows': len(gen_table.rows),
                'ref_rows': len(ref_table.rows),
                'gen_cols': len(gen_table.rows[0].cells),
                'ref_cols': len(ref_table.rows[0].cells),
                'gen_pct': gen_pct,
                'ref_pct': ref_pct,
                'struct_match': struct_match,
                'fill_match': fill_match
            })
    
    # Вывод результатов
    if differences:
        print("❌ НАЙДЕНЫ ОТЛИЧИЯ:\n")
        for diff in differences:
            print(f"Таблица {diff['table']}:")
            if not diff['struct_match']:
                print(f"  ❌ Структура: {diff['gen_rows']}×{diff['gen_cols']} (сгенерирована) " + 
                      f"vs {diff['ref_rows']}×{diff['ref_cols']} (эталон)")
            if not diff['fill_match']:
                print(f"  ❌ Заполнение: {diff['gen_pct']:.1f}% (сгенерирована) " +
                      f"vs {diff['ref_pct']:.1f}% (эталон)")
            print()
        
        return False, differences
    else:
        print("✅ ВСЕ ТАБЛИЦЫ СОВПАДАЮТ!")
        return True, []

if __name__ == '__main__':
    match, differences = analyze_documents()
    if not match:
        print("\n⚠️  Требуется динамическое сканирование и построение маппинга")
