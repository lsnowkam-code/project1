"""
Advanced Dynamic Mapping - Для каждой ячейки определяем источник данных
"""
from docx import Document
from pathlib import Path
import pandas as pd
import csv
from difflib import SequenceMatcher

def similar(a, b):
    return SequenceMatcher(None, a.lower(), b.lower()).ratio()

def load_table_source_mapping():
    mapping = {}
    with open('/workspaces/project/input/mappings/table_source_data_mapping.csv', 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f, delimiter=';')
        for row in reader:
            table_name = row['Таблица'].strip()
            source = row['Источник'].strip()
            mapping[table_name] = source
    return mapping

def get_table_header(table):
    """Извлекает все текст из первых 3 строк"""
    header = []
    for i in range(min(3, len(table.rows))):
        for cell in table.rows[i].cells:
            text = cell.text.strip()
            if text:
                header.append(text)
    return ' '.join(header)

def find_source(table_header, mapping):
    """Находит источник для таблицы"""
    best = None
    best_score = 0.0
    
    for table_name, source in mapping.items():
        score = similar(table_header, table_name)
        if score > best_score:
            best_score = score
            best = (table_name, source)
    
    return best if best_score > 0.4 else (None, None)

print("=" * 100)
print("ДИНАМИЧЕСКОЕ СОЗДАНИЕ МАППИНГА ЯЧЕЕК")
print("=" * 100)
print()

doc = Document('/workspaces/project/output/как должно быть!.docx')
table_mapping = load_table_source_mapping()

all_cells = []

for table_num in range(len(doc.tables)):
    table = doc.tables[table_num]
    
    header = get_table_header(table)
    table_name, source = find_source(header, table_mapping)
    
    print(f"📊 Таблица {table_num + 1}:")
    print(f"   Источник: {source if source else '❌ НЕИЗВЕСТНО'}")
    
    # Определяем структуру
    header_rows = 0
    for i in range(min(5, len(table.rows))):
        row_text = ''.join(
            cell.text.strip() 
            for cell in table.rows[i].cells
        )
        has_numbers = any(
            c.isdigit() for c in row_text if c not in '., (%)'
        )
        if not has_numbers:
            header_rows = i + 1
        else:
            break
    
    print(f"   Заголовков: {header_rows} строк")
    print()
    
    # Для каждой ячейки
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell_text = cell.text.strip()
            
            # Определяем тип ячейки
            if row_idx < header_rows:
                cell_type = 'HEADER'
            elif col_idx == 0:
                cell_type = 'OKVED_CODE'
            elif col_idx == 1 and row_idx >= header_rows:
                cell_type = 'OKVED_NAME'
            elif cell_text == '—':
                cell_type = 'DASH'
            elif cell_text == '':
                cell_type = 'EMPTY'
            else:
                cell_type = 'DATA'
            
            all_cells.append({
                'table_id': table_num + 1,
                'table_source': source if source else 'UNKNOWN',
                'row_id': row_idx,
                'col_id': col_idx,
                'is_header': row_idx < header_rows,
                'cell_type': cell_type,
                'cell_value': cell_text[:80],
                'cell_length': len(cell_text),
                'is_number': any(c.isdigit() for c in cell_text),
                'is_text': cell_text and not any(c.isdigit() for c in cell_text.split())
            })

# Сохраняем маппинг
df = pd.DataFrame(all_cells)
output = '/workspaces/project/output/ДИНАМИЧЕСКИЙ_МАППИНГ_ЯЧЕЕК.csv'
df.to_csv(output, encoding='utf-8-sig', index=False)

print("=" * 100)
print(f"✅ Маппинг создан: {output}")
print(f"   Всего ячеек: {len(all_cells)}")
print()
print("📊 СТАТИСТИКА:")
print()
print("По типам ячеек:")
print(df['cell_type'].value_counts())
print()
print("По таблицам:")
print(df.groupby('table_id')['table_source'].first())
print()
print("По типам данных в ячейках:")
print(f"  - HEADER: {(df['cell_type'] == 'HEADER').sum()}")
print(f"  - DATA: {(df['cell_type'] == 'DATA').sum()}")
print(f"  - DASH: {(df['cell_type'] == 'DASH').sum()}")
print(f"  - EMPTY: {(df['cell_type'] == 'EMPTY').sum()}")
print(f"  - OKVED CODE: {(df['cell_type'] == 'OKVED_CODE').sum()}")
print(f"  - OKVED NAME: {(df['cell_type'] == 'OKVED_NAME').sum()}")
print()
print("=" * 100)
