from docx import Document

generated = Document('/workspaces/project/output/Бюллетень_17.2.8_ГОТОВЫЙ.docx')
reference = Document('/workspaces/project/output/как должно быть!.docx')

# Table 17
print("=" * 80)
print("ТАБЛИЦА 17 - СТРУКТУРА")
print("=" * 80)
t17_gen = generated.tables[16]
t17_ref = reference.tables[16]
print(f"Сгенерирована: {len(t17_gen.rows)} строк, {len(t17_gen.rows[0].cells)} столбцов")
print(f"Эталон:       {len(t17_ref.rows)} строк, {len(t17_ref.rows[0].cells)} столбцов")

# Table 18
print()
print("=" * 80)
print("ТАБЛИЦА 18 - СТРУКТУРА")
print("=" * 80)
t18_gen = generated.tables[17]
t18_ref = reference.tables[17]
print(f"Сгенерирована: {len(t18_gen.rows)} строк, {len(t18_gen.rows[0].cells)} столбцов")
print(f"Эталон:       {len(t18_ref.rows)} строк, {len(t18_ref.rows[0].cells)} столбцов")

print()
print("Заголовки таблицы 18 (сгенерирована):")
for i in range(min(5, len(t18_gen.rows[0].cells))):
    print(f"  [{i}]: {t18_gen.rows[0].cells[i].text[:50]}")

print()
print("Заголовки таблицы 18 (эталон):")
for i in range(min(5, len(t18_ref.rows[0].cells))):
    print(f"  [{i}]: {t18_ref.rows[0].cells[i].text[:50]}")
